"""Offline Excel and PDF reporting for VarianceDesk."""
from __future__ import annotations

from dataclasses import asdict, is_dataclass
from pathlib import Path
from typing import Any, Iterable, Mapping


def _row(value: Any) -> dict[str, Any]:
    """Accept domain dataclasses, mappings, or lightweight record objects."""
    if is_dataclass(value):
        return asdict(value)
    if isinstance(value, Mapping):
        return dict(value)
    return {key: getattr(value, key) for key in
            ("period", "department", "account", "category", "budget", "actual")
            if hasattr(value, key)}


def normalize_rows(records: Iterable[Any]) -> list[dict[str, Any]]:
    rows = [_row(item) for item in records]
    required = {"period", "department", "account", "category", "budget", "actual"}
    for row in rows:
        if "category" not in row and "kind" in row:
            row["category"] = row["kind"]
        missing = required - row.keys()
        if missing:
            raise ValueError(f"Report data is missing fields: {', '.join(sorted(missing))}")
        row["budget"] = float(row["budget"])
        row["actual"] = float(row["actual"])
        row["period"] = str(row["period"])
    return rows


def _variance(row: Mapping[str, Any]) -> float:
    # Revenue: actual above budget is favorable; expenses: actual below budget is favorable.
    raw = float(row["actual"]) - float(row["budget"])
    return raw if str(row["category"]).lower() == "revenue" else -raw


def reconcile(records: Iterable[Any]) -> dict[str, float]:
    rows = normalize_rows(records)
    return {
        "budget": round(sum(r["budget"] for r in rows), 2),
        "actual": round(sum(r["actual"] for r in rows), 2),
        "favorable_variance": round(sum(_variance(r) for r in rows), 2),
    }


def export_excel(records: Iterable[Any], destination: str | Path,
                 forecast: Iterable[Mapping[str, Any]] = (),
                 assumptions: Mapping[str, Any] | None = None) -> Path:
    """Create a self-contained management workbook and return its path."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rows = normalize_rows(records)
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="16324F")

    def sheet(name: str, headers: list[str], values: Iterable[Iterable[Any]]):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for value in values:
            ws.append(list(value))
        ws.freeze_panes = "A2"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(32, max(12, max(len(str(c.value or "")) for c in column) + 2))
        return ws

    totals = reconcile(rows)
    sheet("Executive Summary", ["Metric", "Value"], totals.items())
    headers = ["Period", "Department", "Account", "Category", "Budget", "Actual"]
    values = [[r[h.lower()] for h in headers] for r in rows]
    sheet("Budget", headers[:-1], [v[:-1] for v in values])
    sheet("Actuals", [h for h in headers if h != "Budget"], [[v[0], v[1], v[2], v[3], v[5]] for v in values])
    variance_values = [[r["period"], r["department"], r["account"], r["category"], r["budget"], r["actual"], _variance(r),
                        None if not r["budget"] else _variance(r) / abs(r["budget"])] for r in rows]
    ws = sheet("Variance Analysis", headers + ["Favorable Variance", "Variance %"], variance_values)
    for cell in ws["H"][1:]: cell.number_format = "0.0%"
    forecast_rows = list(forecast)
    sheet("Forecast", list(forecast_rows[0].keys()) if forecast_rows else ["Period", "Base", "Optimistic", "Pessimistic"],
          [r.values() for r in forecast_rows])
    sheet("Scenario Assumptions", ["Assumption", "Value"], (assumptions or {}).items())
    sheet("Definitions", ["Term", "Definition"], [
        ("Favorable variance", "Revenue actual above budget or expense actual below budget."),
        ("Variance percentage", "Favorable variance divided by absolute budget; blank when budget is zero."),
        ("Forecast", "Transparent rolling scenario projection based on configured growth and expense assumptions."),
    ])
    wb.save(path)
    return path


def export_pdf(records: Iterable[Any], destination: str | Path,
               commentary: str | None = None) -> Path:
    """Create a concise offline executive PDF report."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    rows = normalize_rows(records)
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    totals = reconcile(rows)
    styles = getSampleStyleSheet()
    elements = [Paragraph("VarianceDesk Executive Performance Summary", styles["Title"]), Spacer(1, 14)]
    summary = [["Metric", "Amount"], ["Budget", f"${totals['budget']:,.2f}"],
               ["Actual", f"${totals['actual']:,.2f}"], ["Favorable variance", f"${totals['favorable_variance']:,.2f}"]]
    table = Table(summary, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .25, colors.grey),
                               ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("PADDING", (0, 0), (-1, -1), 6)]))
    elements.extend([table, Spacer(1, 16), Paragraph("Management commentary", styles["Heading2"]),
                     Paragraph(commentary or "Commentary is generated from recorded financial results; it does not infer causes.", styles["BodyText"]), Spacer(1, 12)])
    contributors = sorted(rows, key=lambda r: abs(_variance(r)), reverse=True)[:8]
    data = [["Period", "Department", "Account", "Favorable variance"]] + [[r["period"], r["department"], r["account"], f"${_variance(r):,.2f}"] for r in contributors]
    elements.append(Paragraph("Major variance contributors", styles["Heading2"]))
    elements.append(Table(data, repeatRows=1, style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")), ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("PADDING", (0, 0), (-1, -1), 5)])))
    SimpleDocTemplate(str(path), pagesize=letter, title="VarianceDesk Report").build(elements)
    return path
