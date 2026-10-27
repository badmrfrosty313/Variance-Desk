"""Unattended end-to-end verification run inside the packaged application."""
from __future__ import annotations

import json
import os
import traceback
from pathlib import Path
from unittest.mock import patch


def run_packaged_qa(output: str | Path) -> int:
    """Drive production widgets in-process and save screenshots and diagnostics."""
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    log: list[str] = []
    screenshots: list[str] = []
    window = None
    app = None
    try:
        os.environ["VARIANCEDESK_DATA_DIR"] = str(output / "user-data")
        from PySide6.QtWidgets import QApplication, QFileDialog
        from PySide6.QtCore import Qt
        from .ui import DashboardWindow, _get
        from .sample_data import demo_transactions
        from .imports import import_file
        from .finance import variance
        app = QApplication.instance() or QApplication(["VarianceDesk QA"])

        def capture(target, name: str) -> None:
            app.processEvents()
            if target.size().width() < 200 or target.size().height() < 150:
                raise AssertionError(f"Screen {name} is empty or undersized: {target.size()}")
            path = output / f"screen-{name.lower().replace(' ', '-')}.png"
            if not target.grab().save(str(path), "PNG"):
                raise AssertionError(f"Failed to save screenshot: {path}")
            screenshots.append(path.name)
            log.append(f"PASS screenshot {name}: {target.size().width()}x{target.size().height()}")

        window = DashboardWindow()
        window.resize(1360, 900)
        window.show()
        app.processEvents()
        if len(window._records) != 48:
            raise AssertionError(f"Expected 48 demo rows, got {len(window._records)}")
        for index, name in enumerate(window.page_names):
            window.nav_buttons[name].click()
            app.processEvents()
            if window.pages.currentIndex() != index:
                raise AssertionError(f"Navigation failed for {name}")
            capture(window.pages.currentWidget(), name)

        # Navigate via real buttons and edit through the production table cellChanged handler.
        window.nav_buttons["Variance Analysis"].click()
        row = next(i for i in range(window.variance_table.rowCount())
                   if window.variance_table.item(i, 0).text() == "2025-12"
                   and window.variance_table.item(i, 1).text() == "Sales")
        window.variance_table.item(row, 5).setText("5050")
        app.processEvents()
        edited = next(r for r in window._records if r["period"] == "2025-12" and r["department"] == "Sales")
        variance_amount, variance_pct = variance(edited["actual"], edited["budget"], edited["kind"])
        if edited["actual"] != 5050 or variance_amount != -266950:
            raise AssertionError(f"Value/variance mismatch after edit: {edited}")
        if variance_pct is None or abs(variance_pct + 98.14) > .02:
            raise AssertionError(f"Percentage mismatch after edit: {variance_pct}")
        log.append("PASS numerical regression: actual=5050 variance=-266950 percent=-98.1%")

        # Sort and filter, then edit a different record by its displayed stable identity.
        window.variance_table.sortItems(1, Qt.SortOrder.DescendingOrder)
        window.period_filter.setCurrentText("2025-12")
        window.department_filter.setCurrentText("Operations")
        app.processEvents()
        if window.variance_table.rowCount() != 1 or window.variance_table.item(0, 1).text() != "Operations":
            raise AssertionError("Period/department filtering failed")
        window.variance_table.item(0, 5).setText("160000")
        app.processEvents()
        expense = next(r for r in window._records if r["period"] == "2025-12" and r["department"] == "Operations")
        if variance(expense["actual"], expense["budget"], expense["kind"])[0] != 7900:
            raise AssertionError(f"Expense variance sign or stable row identity failed: {expense}")
        window.variance_table.item(0, 4).setText("180000")
        app.processEvents()
        expense = next(r for r in window._records if r["period"] == "2025-12" and r["department"] == "Operations")
        if variance(expense["actual"], expense["budget"], expense["kind"])[0] != 20000:
            raise AssertionError(f"Budget edit did not recalculate expense variance: {expense}")
        window.period_filter.setCurrentText("All periods")
        window.department_filter.setCurrentText("All departments")
        app.processEvents()
        log.append("PASS sorted/filter UI editing and expense/budget recalculation")

        # Exercise assumptions, scenarios, import dialogs, and report dialogs through buttons.
        window.nav_buttons["Forecasting"].click()
        first_base = next(r for r in window._forecast if r["scenario"] == "Base")
        window.revenue_growth.setValue(9.0)
        window.recalculate_forecast_button.click()
        app.processEvents()
        changed_base = next(r for r in window._forecast if r["scenario"] == "Base")
        if changed_base["revenue"] == first_base["revenue"]:
            raise AssertionError("Forecast did not respond to revenue growth control")
        window.scenario_filter.setCurrentText("Optimistic")
        app.processEvents()
        if window.forecast_table.rowCount() != 12:
            raise AssertionError("Scenario control did not filter forecast rows")
        log.append("PASS forecast controls and scenario selection")

        csv_path = output / "qa-import.csv"
        csv_path.write_text("period,department,account,kind,budget,actual\n2026-01,QA,QA Revenue,revenue,200,225\n", encoding="utf-8")
        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["period", "department", "account", "kind", "budget", "actual"])
        sheet.append(["2026-01", "QA", "QA Expense", "expense", 80, 75])
        xlsx_path = output / "qa-import.xlsx"; workbook.save(xlsx_path)
        window.nav_buttons["Data Management"].click()
        with patch.object(QFileDialog, "getOpenFileName", return_value=(str(csv_path), "CSV")):
            window.import_button.click(); app.processEvents()
        with patch.object(QFileDialog, "getOpenFileName", return_value=(str(xlsx_path), "XLSX")):
            window.import_button.click(); app.processEvents()
        if not any(_get(r, "account") == "QA Revenue" for r in window._records) or not any(_get(r, "account") == "QA Expense" for r in window._records):
            raise AssertionError("CSV/XLSX UI import did not update the application dataset")
        window.nav_buttons["Reports"].click()
        excel_path, pdf_path = output / "qa-report.xlsx", output / "qa-report.pdf"
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(excel_path), "Excel")):
            window.excel_button.click(); app.processEvents()
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(pdf_path), "PDF")):
            window.pdf_button.click(); app.processEvents()
        if excel_path.stat().st_size <= 0 or pdf_path.stat().st_size <= 0:
            raise AssertionError("UI export produced an empty report")
        from openpyxl import load_workbook
        exported = load_workbook(excel_path, data_only=True)
        actuals = {exported["Actuals"].cell(i, 3).value: exported["Actuals"].cell(i, 5).value for i in range(2, exported["Actuals"].max_row + 1)}
        if actuals.get("Subscription Revenue") != 5050:
            raise AssertionError("Excel report does not contain the edited actual")
        log.append("PASS CSV/XLSX import and reconciled Excel/PDF exports")

        window.close(); window.store.close(); app.processEvents()
        reopened = DashboardWindow()
        reopened.show(); app.processEvents()
        persisted = {r["account"]: r for r in reopened._records}
        if persisted["Subscription Revenue"]["actual"] != 5050:
            raise AssertionError("SQLite edit was not present after application reopen")
        if "QA Expense" not in persisted:
            raise AssertionError("Imported XLSX data did not persist after reopen")
        reopened.close(); reopened.store.close(); app.processEvents()
        log.append("PASS close/reopen SQLite persistence")
        result = {"status": "passed", "screenshots": screenshots, "checks": log}
        (output / "qa-report.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        (output / "qa.log").write_text("\n".join(log) + "\n", encoding="utf-8")
        return 0
    except Exception:
        log.append(traceback.format_exc())
        if window is not None:
            try:
                path = output / "failure-screen.png"
                window.grab().save(str(path), "PNG")
                screenshots.append(path.name)
            except Exception:
                pass
        (output / "qa.log").write_text("\n".join(log) + "\n", encoding="utf-8")
        (output / "qa-report.json").write_text(json.dumps({"status": "failed", "screenshots": screenshots, "checks": log}, indent=2), encoding="utf-8")
        return 1
    finally:
        if window is not None:
            try: window.close(); window.store.close()
            except Exception: pass
