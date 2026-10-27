from pathlib import Path
from variancedesk.reporting import export_excel, export_pdf, reconcile


ROWS = [
    {"period": "2025-01", "department": "Sales", "account": "Subscription Revenue", "category": "Revenue", "budget": 100.0, "actual": 115.0},
    {"period": "2025-01", "department": "Operations", "account": "Hosting", "category": "Expense", "budget": 50.0, "actual": 42.0},
]


def test_reconciliation_uses_financial_favorability():
    assert reconcile(ROWS) == {"budget": 150.0, "actual": 157.0, "favorable_variance": 23.0}


def test_excel_report_contains_required_sheets(tmp_path: Path):
    from openpyxl import load_workbook
    out = export_excel(ROWS, tmp_path / "report.xlsx")
    wb = load_workbook(out, data_only=True)
    assert {"Executive Summary", "Budget", "Actuals", "Variance Analysis", "Forecast", "Scenario Assumptions", "Definitions"} <= set(wb.sheetnames)
    assert wb["Executive Summary"]["B2"].value == 150


def test_pdf_report_is_generated(tmp_path: Path):
    out = export_pdf(ROWS, tmp_path / "report.pdf")
    assert out.exists() and out.stat().st_size > 1000
