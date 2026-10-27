from __future__ import annotations

from variancedesk.__main__ import main, self_test
from variancedesk.sample_data import demo_transactions
from variancedesk.finance import ForecastAssumptions, build_forecast
from variancedesk.ui import DashboardWindow
from PySide6.QtWidgets import QFileDialog
from PySide6.QtWidgets import QMessageBox


def test_dashboard_loads_demo_records_and_filters(qtbot, monkeypatch, tmp_path):
    monkeypatch.setenv("VARIANCEDESK_DATA_DIR", str(tmp_path))
    window = DashboardWindow(demo_transactions())
    qtbot.addWidget(window)
    assert window.table.rowCount() == 48
    assert "$" in window.cards["revenue"].value.text()
    assert window.table.editTriggers() == window.table.EditTrigger.NoEditTriggers
    window.department_filter.setCurrentText("Sales")
    assert window.table.rowCount() == 24


def test_application_headless_diagnostic():
    assert self_test() == 0
    assert main(["--self-test", "--headless"]) == 0


def test_gui_import_forecast_exports_and_persistence(qtbot, monkeypatch, tmp_path):
    monkeypatch.setenv("VARIANCEDESK_DATA_DIR", str(tmp_path))
    source = tmp_path / "financials.csv"
    source.write_text("period,department,account,kind,budget,actual\n2025-01,Sales,Subscriptions,revenue,100,120\n2025-01,Ops,Hosting,expense,50,40\n", encoding="utf-8")
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: (str(source), "CSV"))
    window = DashboardWindow()
    qtbot.addWidget(window)
    window.show_page(3)
    window.import_button.click()
    assert window.table.rowCount() == 2
    assert "Imported 2" in window.import_status.text()

    window.show_page(2)
    base = [r["revenue"] for r in window._forecast if r["scenario"] == "Base"]
    base_expenses = [r["expenses"] for r in window._forecast if r["scenario"] == "Base"]
    assert window.revenue_growth.value() == 3
    assert window._assumption_dict()["revenue_growth"] == .03
    window.revenue_growth.setValue(10)
    window.recalculate_forecast_button.click()
    changed = [r["revenue"] for r in window._forecast if r["scenario"] == "Base"]
    assert changed != base
    assert changed[0] == 120 * 1.10
    window.scenario_filter.setCurrentText("Base")
    visible = {window.forecast_table.item(i, 0).text(): window.forecast_table.item(i, 2).text() for i in range(window.forecast_table.rowCount())}
    assert visible[window._forecast[0]["period"]] == f"${changed[0]:,.0f}"
    window.expense_growth.setValue(8)
    window.recalculate_forecast_button.click()
    expenses_changed = [r["expenses"] for r in window._forecast if r["scenario"] == "Base"]
    assert expenses_changed != base_expenses
    assert expenses_changed[0] == 40 * 1.08
    repeated = list(window._forecast)
    window.recalculate_forecast_button.click()
    assert window._forecast == repeated
    window.scenario_filter.setCurrentText("Optimistic")
    assert window.forecast_table.rowCount() == 12
    assert window.forecast_table.item(0, 1).text() == "Optimistic"
    optimistic = window.forecast_table.item(0, 2).text()
    window.scenario_filter.setCurrentText("Pessimistic")
    assert window.forecast_table.item(0, 1).text() == "Pessimistic"
    assert window.forecast_table.item(0, 2).text() != optimistic
    expected = build_forecast(window._monthly_history(window._records), ForecastAssumptions(**window._assumption_dict()))
    first = next(r for r in expected if r["scenario"] == "Base")
    assert first["revenue"] == window._forecast[0]["revenue"]
    assert window._monthly_history(window._records)[-1]["actual_revenue"] == 120

    xlsx, pdf = tmp_path / "gui.xlsx", tmp_path / "gui.pdf"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(xlsx), "XLSX"))
    window.show_page(4); window.excel_button.click()
    assert xlsx.exists() and xlsx.stat().st_size > 0
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(pdf), "PDF"))
    window.pdf_button.click()
    assert pdf.exists() and pdf.stat().st_size > 0
    window.close(); window.store.close()

    reopened = DashboardWindow()
    qtbot.addWidget(reopened)
    assert not reopened.demo_mode
    assert reopened.table.rowCount() == 2
    reopened.store.close()


def test_editing_values_updates_all_views_and_persists(qtbot, monkeypatch, tmp_path):
    monkeypatch.setenv("VARIANCEDESK_DATA_DIR", str(tmp_path))
    records = [
        {"period": "2025-12", "department": "Sales", "account": "Subscription Revenue", "kind": "revenue", "budget": 272000.0, "actual": 274720.0},
        {"period": "2025-12", "department": "Operations", "account": "Operating Expenses", "kind": "expense", "budget": 100000.0, "actual": 102000.0},
    ]
    window = DashboardWindow(records)
    qtbot.addWidget(window)
    window.show_page(1)
    findrow = lambda dept: next(i for i in range(window.variance_table.rowCount()) if window.variance_table.item(i, 1).text() == dept)

    revenue_row = findrow("Sales")
    window.variance_table.item(revenue_row, 5).setText("5050")
    assert window.variance_table.item(findrow("Sales"), 6).text() == "($266,950)"
    assert window.variance_table.item(findrow("Sales"), 7).text() == "-98.1%"
    assert window.cards["revenue"].value.text() == "$5,050"
    assert window._monthly_history(window._records)[-1]["actual_revenue"] == 5050

    expense_row = findrow("Operations")
    window.variance_table.item(expense_row, 5).setText("90000")
    assert window.variance_table.item(findrow("Operations"), 6).text() == "$10,000"
    expense_row = findrow("Operations")
    window.variance_table.item(expense_row, 4).setText("95000")
    assert window.variance_table.item(findrow("Operations"), 6).text() == "$5,000"

    window.show_page(2)
    assert window._forecast[0]["revenue"] == 5050 * 1.03
    window.show_page(4)
    xlsx, pdf = tmp_path / "edited.xlsx", tmp_path / "edited.pdf"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(xlsx), "XLSX"))
    window.excel_button.click()
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(pdf), "PDF"))
    window.pdf_button.click()
    from openpyxl import load_workbook
    book = load_workbook(xlsx, data_only=True)
    exported = {book["Actuals"].cell(r, 3).value: book["Actuals"].cell(r, 5).value for r in range(2, book["Actuals"].max_row + 1)}
    assert exported["Subscription Revenue"] == 5050
    assert pdf.exists() and pdf.stat().st_size > 0
    window.store.close()

    reopened = DashboardWindow()
    qtbot.addWidget(reopened)
    assert next(r for r in reopened._records if r["department"] == "Sales")["actual"] == 5050
    assert next(r for r in reopened._records if r["department"] == "Operations")["budget"] == 95000
    reopened.store.close()


def test_invalid_financial_edit_is_rejected(qtbot, monkeypatch, tmp_path):
    monkeypatch.setenv("VARIANCEDESK_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: None)
    rows = [{"period": "2025-12", "department": "Sales", "account": "Revenue", "kind": "revenue", "budget": 100, "actual": 110}]
    window = DashboardWindow(rows); qtbot.addWidget(window); window.show_page(1)
    window.variance_table.item(0, 5).setText("not money")
    assert window.variance_table.item(0, 5).text() == "$110"
    assert window._records[0]["actual"] == 110
    assert window.store.fetch() == []
    window.store.close()
